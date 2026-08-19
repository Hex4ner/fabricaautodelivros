#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Escrita na planilha mestre por codigo, nunca por leitura.

Uso tipico dentro de um script de autor:

    from lote_helper import Planilha
    p = Planilha("saida/CITACOES_207_v03.xlsx")
    p.citacao(...); p.auditoria(...); p.apocrifa(...); p.placar(...)
    p.salvar()

A planilha e sempre aberta, acrescida e salva. Nunca recriada.
"""
import json
import os
import shutil
from datetime import date

HOJE = str(date.today())


class Planilha:
    def __init__(self, caminho, base=None):
        import openpyxl
        if not os.path.exists(caminho):
            if not base:
                raise SystemExit(f"FALHA: {caminho} nao existe e nenhuma base foi dada")
            shutil.copyfile(base, caminho)
            print(f"  [novo] {caminho} criado a partir de {base}")
        self.caminho = caminho
        self.wb = openpyxl.load_workbook(caminho)
        self.cit = self.wb["CITACOES"]
        self.aud = self.wb["AUDITORIA_96"]
        self.apo = self.wb["APOCRIFAS"]
        self.pla = self.wb["PLACAR"]

    # ---------- aba CITACOES ----------
    def citacao(self, num, nome, ordem, corpo, rodape="", original="", idioma="",
                obra="", locus="", edicao="", selo="", podio="Não", nivel="",
                url="", conf="", dominio="Sim", obs=""):
        self.cit.append([str(num), nome, str(ordem), corpo, rodape, original, idioma,
                         obra, locus, edicao, selo, podio, nivel, url,
                         str(conf) if conf else "", dominio, HOJE, obs])

    def vaga(self, num, nome, ordem, obs, dominio="Sim"):
        self.cit.append([str(num), nome, str(ordem), "(vaga aberta)", None, None, None,
                         None, None, None, None, "Não", None, None, None, dominio, HOJE, obs])

    # ---------- aba AUDITORIA_96 ----------
    def auditoria(self, num, veredito, selo, prova, substituta, prioridade):
        """Atualiza a linha existente do expoente (a frase antiga ja esta la)."""
        alvo = str(num).strip()
        for row in self.aud.iter_rows(min_row=2):
            if str(row[0].value).strip() == alvo:
                row[3].value = veredito
                row[4].value = selo
                row[5].value = prova
                row[6].value = substituta
                row[7].value = prioridade
                row[8].value = HOJE
                return True
        raise SystemExit(f"FALHA: expoente {num} nao encontrado na AUDITORIA_96")

    # ---------- aba APOCRIFAS ----------
    def apocrifa(self, num, autor, frase, origem, primeira, padroes, defesa):
        self.apo.append([str(num), autor, frase, origem, primeira, padroes, defesa])

    # ---------- aba PLACAR ----------
    def placar(self, num, fechadas, status, lote):
        alvo = str(num).strip()
        for row in self.pla.iter_rows(min_row=2):
            if str(row[0].value).strip() == alvo:
                row[7].value = fechadas
                row[8].value = status
                row[9].value = lote
                return True
        raise SystemExit(f"FALHA: expoente {num} nao encontrado no PLACAR")

    def salvar(self):
        self.wb.save(self.caminho)
        print(f"  [gravado] {self.caminho}")


def registrar(chave, autor, fechadas, meta, status, **extra):
    """Acrescenta uma linha ao indice de retomada. Nunca reescreve o arquivo."""
    linha = {"chave": chave, "autor": autor, "data": HOJE, "modo": "AUDITORIA+BUSCA",
             "fechadas": fechadas, "meta": meta, "status": status}
    linha.update(extra)
    with open("_registro-citacoes.jsonl", "a", encoding="utf-8") as f:
        f.write(json.dumps(linha, ensure_ascii=False) + "\n")
    print(f"  [registro] {autor}: {fechadas}/{meta} {status}")

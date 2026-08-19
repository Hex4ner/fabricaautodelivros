#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
trava.py - checagem mecânica da planilha de citações do projeto IAP.

Existe por um motivo: um agente autônomo verificando centenas de citações vai, em algum momento,
escrever um locus que parece certo sem ter recuperado nada. Nenhuma regra escrita em prosa impede
isso. Uma condição checável, sim.

A regra-mãe: selo forte (PRIMÁRIA ou CONTEMPORÂNEA) exige endereço de prova recuperado.
Sem URL, sem selo forte.

Uso:
    python scripts/trava.py saida/CITACOES_207_v03.xlsx

Saída: relatório em texto. Código de saída 1 se houver qualquer FALHA.
"""

import sys
import unicodedata
from collections import Counter

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl nao encontrado. Instale com: pip install openpyxl")

SELOS_FORTES = {"PRIMÁRIA", "CONTEMPORÂNEA"}
SELOS_VALIDOS = SELOS_FORTES | {"ATRIBUÍDA", "DISPUTADA", "APÓCRIFA"}
NIVEIS_FORTES = {"N1", "N2", "N3"}
VAGA = "(vaga aberta)"


def norm(s):
    """Maiuscula, sem acento, sem espaco nas pontas. Para comparar selo digitado torto."""
    s = str(s or "").strip().upper()
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


NORM_FORTES = {norm(x) for x in SELOS_FORTES}
NORM_VALIDOS = {norm(x) for x in SELOS_VALIDOS}


def carregar(caminho):
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    if "CITACOES" not in wb.sheetnames:
        sys.exit(f"FALHA: aba CITACOES nao encontrada em {caminho}")
    ws = wb["CITACOES"]
    linhas = list(ws.iter_rows(values_only=True))
    wb.close()
    if not linhas:
        sys.exit("FALHA: aba CITACOES vazia")
    cab = [str(c).strip() if c is not None else "" for c in linhas[0]]
    return cab, linhas[1:]


def coluna(cab, alvo):
    """Acha a coluna pelo inicio do cabecalho, para tolerar renomeacao parcial."""
    a = norm(alvo)
    for i, nome in enumerate(cab):
        if norm(nome).startswith(a):
            return i
    sys.exit(f"FALHA: coluna '{alvo}' nao encontrada. Cabecalho lido: {cab}")


def main():
    if len(sys.argv) < 2:
        sys.exit("Uso: python scripts/trava.py <caminho-da-planilha.xlsx>")
    caminho = sys.argv[1]
    cab, dados = carregar(caminho)

    C = {k: coluna(cab, k) for k in
         ["Nº", "Nome", "Ordem", "CORPO", "RODAPÉ", "Texto original", "Idioma",
          "Locus", "Selo", "Pódio", "Nível", "Fonte recuperada", "Conf", "Data"]}

    falhas, avisos = [], []
    vistos = Counter()
    fortes = preenchidas = vagas = 0

    for n, r in enumerate(dados, start=2):
        def v(k):
            i = C[k]
            return str(r[i]).strip() if i < len(r) and r[i] is not None else ""

        num, nome, ordem = v("Nº"), v("Nome"), v("Ordem")
        corpo, selo, url = v("CORPO"), v("Selo"), v("Fonte recuperada")
        locus, nivel, podio = v("Locus"), v("Nível"), v("Pódio")
        original, idioma, rodape = v("Texto original"), v("Idioma"), v("RODAPÉ")
        conf = v("Conf")
        s = norm(selo)

        vistos[(num, ordem)] += 1
        if corpo.startswith("(vaga"):
            vagas += 1
            if selo:
                falhas.append(f"L{n} {nome} #{ordem}: vaga aberta nao pode ter selo ('{selo}')")
            continue
        if not corpo:
            avisos.append(f"L{n} {nome} #{ordem}: CORPO vazio e sem marca de vaga aberta")
            continue

        preenchidas += 1

        if not selo:
            falhas.append(f"L{n} {nome} #{ordem}: citacao preenchida sem selo")
        elif s not in NORM_VALIDOS:
            falhas.append(f"L{n} {nome} #{ordem}: selo invalido ('{selo}')")

        if s in NORM_FORTES:
            fortes += 1
            # A REGRA-MAE
            if not url:
                falhas.append(f"L{n} {nome} #{ordem}: selo {selo} SEM URL de prova recuperada")
            if not locus:
                falhas.append(f"L{n} {nome} #{ordem}: selo {selo} sem locus")
            if nivel and nivel.upper() not in NIVEIS_FORTES:
                falhas.append(f"L{n} {nome} #{ordem}: selo {selo} com nivel {nivel} (exige N1-N3)")
            if norm(podio) != "SIM":
                falhas.append(f"L{n} {nome} #{ordem}: selo forte mas Podio='{podio}'")
        else:
            if norm(podio) == "SIM":
                falhas.append(f"L{n} {nome} #{ordem}: Podio=Sim com selo fraco ('{selo}')")

        # traducao declarada: idioma estrangeiro exige a formula no rodape
        if idioma and not norm(idioma).startswith("PORTUGU"):
            if rodape and "TRADUCAO DO" not in norm(rodape):
                falhas.append(f"L{n} {nome} #{ordem}: rodape sem declaracao de traducao")
            if not original:
                avisos.append(f"L{n} {nome} #{ordem}: idioma {idioma} sem texto original recuperado")

        if conf:
            try:
                if float(str(conf).replace("%", "").replace(",", ".")) < 95:
                    avisos.append(f"L{n} {nome} #{ordem}: confiabilidade {conf} abaixo do piso de 95")
            except ValueError:
                avisos.append(f"L{n} {nome} #{ordem}: confiabilidade ilegivel ('{conf}')")

    for (num, ordem), q in vistos.items():
        if q > 1:
            falhas.append(f"Duplicidade: expoente {num} ordem {ordem} aparece {q} vezes")

    print(f"\nTRAVA MECANICA - {caminho}")
    print(f"  linhas: {len(dados)} | citacoes preenchidas: {preenchidas} "
          f"| selo forte: {fortes} | vagas abertas: {vagas}")
    print(f"  FALHAS: {len(falhas)} | AVISOS: {len(avisos)}\n")

    for f in falhas:
        print(f"  FALHA  {f}")
    for a in avisos:
        print(f"  aviso  {a}")

    if falhas:
        print(f"\nRESULTADO: REPROVADO. Corrija as {len(falhas)} falhas antes de fechar o lote.")
        return 1
    print("\nRESULTADO: APROVADO. Nenhuma citacao com selo forte sem prova recuperada.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
